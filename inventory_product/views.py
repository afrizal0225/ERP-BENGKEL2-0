from django.shortcuts import render, get_object_or_404, redirect
from .models import PenerimaanBarang, PenerimaanBarangDetail, KeluarBarang, KeluarBarangDetail, InventoryProduct, PenerimaanRawMaterial, PenerimaanRawMaterialDetail, KeluarRawMaterial, KeluarRawMaterialDetail, InventoryRawMaterial
from product.models import Product, RawMaterial
from django.core.exceptions import ObjectDoesNotExist
from django.contrib import messages
import openpyxl
from datetime import datetime
from openpyxl import Workbook
from django.http import HttpResponse
from django import forms

class BulkUploadForm(forms.Form):
    file = forms.FileField()

# Penerimaan Barang Views
def penerimaan_list(request):
    penerimaans = PenerimaanBarang.objects.all()
    return render(request, 'inventory_product/penerimaan_list.html', {'penerimaans': penerimaans})

def penerimaan_create(request):
    if request.method == 'POST':
        id_penerimaan = request.POST.get('id_penerimaan')
        id_supplier = request.POST.get('id_supplier')
        tanggal_terima = request.POST.get('tanggal_terima')
        keterangan = request.POST.get('keterangan')
        PenerimaanBarang.objects.create(
            id_penerimaan=id_penerimaan,
            id_supplier=id_supplier,
            tanggal_terima=tanggal_terima,
            keterangan=keterangan
        )
        messages.success(request, 'Penerimaan created successfully.')
        return redirect('penerimaan_list')
    return render(request, 'inventory_product/penerimaan_form.html', {'title': 'Create Penerimaan'})

def penerimaan_update(request, pk):
    penerimaan = get_object_or_404(PenerimaanBarang, pk=pk)
    if request.method == 'POST':
        penerimaan.id_supplier = request.POST.get('id_supplier')
        penerimaan.tanggal_terima = request.POST.get('tanggal_terima')
        penerimaan.keterangan = request.POST.get('keterangan')
        penerimaan.save()
        messages.success(request, 'Penerimaan updated successfully.')
        return redirect('penerimaan_list')
    return render(request, 'inventory_product/penerimaan_form.html', {'penerimaan': penerimaan, 'title': 'Update Penerimaan'})

def penerimaan_delete(request, pk):
    penerimaan = get_object_or_404(PenerimaanBarang, pk=pk)
    if request.method == 'POST':
        penerimaan.delete()
        messages.success(request, 'Penerimaan deleted successfully.')
        return redirect('penerimaan_list')
    return render(request, 'inventory_product/penerimaan_confirm_delete.html', {'penerimaan': penerimaan})

def penerimaan_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 10:
                        continue
                    id_penerimaan, tanggal_terima_str, Supplier, BRAND, sku_product, Nama_product, quantity, satuan, harga_satuan, kategori = row[:10]
                    try:
                        tanggal_terima = datetime.strptime(tanggal_terima_str, '%Y-%m-%d').date()
                        try:
                            product = Product.objects.get(sku=sku_product)
                            if product.name != Nama_product:
                                messages.warning(request, f"Product name '{Nama_product}' does not match database name '{product.name}' for SKU {sku_product}, skipping row.")
                                continue
                        except ObjectDoesNotExist:
                            messages.warning(request, f"SKU {sku_product} not found in products, skipping row.")
                            continue
                        penerimaan, created = PenerimaanBarang.objects.get_or_create(
                            id_penerimaan=id_penerimaan,
                            defaults={'id_supplier': Supplier, 'tanggal_terima': tanggal_terima, 'keterangan': ''}
                        )
                        PenerimaanBarangDetail.objects.create(
                            id_penerimaan=penerimaan,
                            tanggal_terima=tanggal_terima,
                            Supplier=Supplier,
                            BRAND=BRAND,
                            id_product=product,
                            Nama_product=Nama_product,
                            quantity=int(quantity),
                            satuan=satuan,
                            harga_satuan=float(harga_satuan),
                            harga_total=int(quantity) * float(harga_satuan),
                            kategori=kategori
                        )
                        created_count += 1
                    except Exception as e:
                        messages.error(request, f"Error processing row for {id_penerimaan}: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} details created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('penerimaan_list')
    else:
        form = BulkUploadForm()
    return render(request, 'inventory_product/penerimaan_bulk_form.html', {'form': form, 'title': 'Bulk Upload Penerimaan Details'})

# Penerimaan Detail Views
def penerimaan_detail_list(request, pk):
    penerimaan = get_object_or_404(PenerimaanBarang, pk=pk)
    details = PenerimaanBarangDetail.objects.filter(id_penerimaan=penerimaan)
    return render(request, 'inventory_product/penerimaan_detail_list.html', {'penerimaan': penerimaan, 'details': details})

def penerimaan_detail_create(request, pk):
    penerimaan = get_object_or_404(PenerimaanBarang, pk=pk)
    if request.method == 'POST':
        # Parse form data
        tanggal_terima = request.POST.get('tanggal_terima')
        Supplier = request.POST.get('Supplier')
        BRAND = request.POST.get('BRAND')
        id_product_id = request.POST.get('id_product')
        Nama_product = request.POST.get('Nama_product')
        quantity = int(request.POST.get('quantity'))
        satuan = request.POST.get('satuan')
        harga_satuan = float(request.POST.get('harga_satuan'))
        harga_total = quantity * harga_satuan
        kategori = request.POST.get('kategori')
        product = get_object_or_404(Product, pk=id_product_id)
        PenerimaanBarangDetail.objects.create(
            id_penerimaan=penerimaan,
            tanggal_terima=tanggal_terima,
            Supplier=Supplier,
            BRAND=BRAND,
            id_product=product,
            Nama_product=Nama_product,
            quantity=quantity,
            satuan=satuan,
            harga_satuan=harga_satuan,
            harga_total=harga_total,
            kategori=kategori
        )
        messages.success(request, 'Detail added successfully.')
        return redirect('penerimaan_detail_list', pk=pk)
    products = Product.objects.all()
    return render(request, 'inventory_product/penerimaan_detail_form.html', {'penerimaan': penerimaan, 'products': products, 'title': 'Add Detail'})

# Similar for KeluarBarang and KeluarBarangDetail

def keluar_list(request):
    keluars = KeluarBarang.objects.all()
    return render(request, 'inventory_product/keluar_list.html', {'keluars': keluars})

def keluar_create(request):
    if request.method == 'POST':
        id_keluar = request.POST.get('id_keluar')
        tanggal_keluar = request.POST.get('tanggal_keluar')
        keterangan = request.POST.get('keterangan')
        KeluarBarang.objects.create(
            id_keluar=id_keluar,
            tanggal_keluar=tanggal_keluar,
            keterangan=keterangan
        )
        messages.success(request, 'Keluar created successfully.')
        return redirect('keluar_list')
    return render(request, 'inventory_product/keluar_form.html', {'title': 'Create Keluar'})

def keluar_update(request, pk):
    keluar = get_object_or_404(KeluarBarang, pk=pk)
    if request.method == 'POST':
        keluar.tanggal_keluar = request.POST.get('tanggal_keluar')
        keluar.keterangan = request.POST.get('keterangan')
        keluar.save()
        messages.success(request, 'Keluar updated successfully.')
        return redirect('keluar_list')
    return render(request, 'inventory_product/keluar_form.html', {'keluar': keluar, 'title': 'Update Keluar'})

def keluar_delete(request, pk):
    keluar = get_object_or_404(KeluarBarang, pk=pk)
    if request.method == 'POST':
        keluar.delete()
        messages.success(request, 'Keluar deleted successfully.')
        return redirect('keluar_list')
    return render(request, 'inventory_product/keluar_confirm_delete.html', {'keluar': keluar})

def keluar_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 8:
                        continue
                    id_keluar, tanggal_keluar_str, BRAND, sku_product, Nama_product, quantity, harga_satuan, kategori = row[:8]
                    try:
                        tanggal_keluar = datetime.strptime(tanggal_keluar_str, '%Y-%m-%d').date()
                        try:
                            product = Product.objects.get(sku=sku_product)
                            if product.name != Nama_product:
                                messages.warning(request, f"Product name '{Nama_product}' does not match database name '{product.name}' for SKU {sku_product}, skipping row.")
                                continue
                        except ObjectDoesNotExist:
                            messages.warning(request, f"SKU {sku_product} not found in products, skipping row.")
                            continue
                        keluar, created = KeluarBarang.objects.get_or_create(
                            id_keluar=id_keluar,
                            defaults={'tanggal_keluar': tanggal_keluar, 'keterangan': ''}
                        )
                        KeluarBarangDetail.objects.create(
                            id_keluar=keluar,
                            tanggal_keluar=tanggal_keluar,
                            BRAND=BRAND,
                            id_product=product,
                            Nama_product=Nama_product,
                            quantity=int(quantity),
                            harga_satuan=float(harga_satuan),
                            harga_total=int(quantity) * float(harga_satuan),
                            kategori=kategori
                        )
                        created_count += 1
                    except Exception as e:
                        messages.error(request, f"Error processing row for {id_keluar}: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} details created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('keluar_list')
    else:
        form = BulkUploadForm()
    return render(request, 'inventory_product/keluar_bulk_form.html', {'form': form, 'title': 'Bulk Upload Keluar Details'})

def keluar_detail_list(request, pk):
    keluar = get_object_or_404(KeluarBarang, pk=pk)
    details = KeluarBarangDetail.objects.filter(id_keluar=keluar)
    return render(request, 'inventory_product/keluar_detail_list.html', {'keluar': keluar, 'details': details})

def keluar_detail_create(request, pk):
    keluar = get_object_or_404(KeluarBarang, pk=pk)
    if request.method == 'POST':
        tanggal_keluar = request.POST.get('tanggal_keluar')
        BRAND = request.POST.get('BRAND')
        id_product_id = request.POST.get('id_product')
        Nama_product = request.POST.get('Nama_product')
        quantity = int(request.POST.get('quantity'))
        harga_satuan = float(request.POST.get('harga_satuan'))
        harga_total = quantity * harga_satuan
        kategori = request.POST.get('kategori')
        product = get_object_or_404(Product, pk=id_product_id)
        KeluarBarangDetail.objects.create(
            id_keluar=keluar,
            tanggal_keluar=tanggal_keluar,
            BRAND=BRAND,
            id_product=product,
            Nama_product=Nama_product,
            quantity=quantity,
            harga_satuan=harga_satuan,
            harga_total=harga_total,
            kategori=kategori
        )
        messages.success(request, 'Detail added successfully.')
        return redirect('keluar_detail_list', pk=pk)
    products = Product.objects.all()
    return render(request, 'inventory_product/keluar_detail_form.html', {'keluar': keluar, 'products': products, 'title': 'Add Detail'})

def inventory_product_list(request):
    from django.db.models import Sum
    products = Product.objects.all()
    inventory_data = []
    for product in products:
        total_in = PenerimaanBarangDetail.objects.filter(id_product=product).aggregate(total=Sum('quantity'))['total'] or 0
        total_out = KeluarBarangDetail.objects.filter(id_product=product).aggregate(total=Sum('quantity'))['total'] or 0
        current_stock = total_in - total_out
        inventory_data.append({
            'product': product,
            'current_stock': current_stock
        })
    return render(request, 'inventory_product/inventory_product_list.html', {'inventories': inventory_data})

def inventory_rawmaterial_list(request):
    from django.db.models import Sum
    rawmaterials = RawMaterial.objects.all()
    inventory_data = []
    for rawmaterial in rawmaterials:
        total_in = PenerimaanRawMaterialDetail.objects.filter(id_rawmaterial=rawmaterial).aggregate(total=Sum('quantity'))['total'] or 0
        total_out = KeluarRawMaterialDetail.objects.filter(id_rawmaterial=rawmaterial).aggregate(total=Sum('quantity'))['total'] or 0
        current_stock = total_in - total_out
        inventory_data.append({
            'rawmaterial': rawmaterial,
            'current_stock': current_stock
        })
    return render(request, 'inventory_product/inventory_rawmaterial_list.html', {'inventories': inventory_data})


# Raw Material Views
def penerimaan_rawmaterial_list(request):
    penerimaans = PenerimaanRawMaterial.objects.all()
    return render(request, 'inventory_product/penerimaan_rawmaterial_list.html', {'penerimaans': penerimaans})

def penerimaan_rawmaterial_create(request):
    if request.method == 'POST':
        id_penerimaan = request.POST.get('id_penerimaan')
        id_supplier = request.POST.get('id_supplier')
        tanggal_terima = request.POST.get('tanggal_terima')
        keterangan = request.POST.get('keterangan')
        PenerimaanRawMaterial.objects.create(
            id_penerimaan=id_penerimaan,
            id_supplier=id_supplier,
            tanggal_terima=tanggal_terima,
            keterangan=keterangan
        )
        messages.success(request, 'Penerimaan Raw Material created successfully.')
        return redirect('penerimaan_rawmaterial_list')
    return render(request, 'inventory_product/penerimaan_rawmaterial_form.html', {'title': 'Create Penerimaan Raw Material'})

def penerimaan_rawmaterial_update(request, pk):
    penerimaan = get_object_or_404(PenerimaanRawMaterial, pk=pk)
    if request.method == 'POST':
        penerimaan.id_supplier = request.POST.get('id_supplier')
        penerimaan.tanggal_terima = request.POST.get('tanggal_terima')
        penerimaan.keterangan = request.POST.get('keterangan')
        penerimaan.save()
        messages.success(request, 'Penerimaan Raw Material updated successfully.')
        return redirect('penerimaan_rawmaterial_list')
    return render(request, 'inventory_product/penerimaan_rawmaterial_form.html', {'penerimaan': penerimaan, 'title': 'Update Penerimaan Raw Material'})

def penerimaan_rawmaterial_delete(request, pk):
    penerimaan = get_object_or_404(PenerimaanRawMaterial, pk=pk)
    if request.method == 'POST':
        penerimaan.delete()
        messages.success(request, 'Penerimaan Raw Material deleted successfully.')
        return redirect('penerimaan_rawmaterial_list')
    return render(request, 'inventory_product/penerimaan_rawmaterial_confirm_delete.html', {'penerimaan': penerimaan})

def penerimaan_rawmaterial_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 10:
                        continue
                    id_penerimaan, tanggal_terima_str, Supplier, BRAND, sku_rawmaterial, Nama_rawmaterial, quantity, satuan, harga_satuan, kategori = row[:10]
                    try:
                        tanggal_terima = datetime.strptime(tanggal_terima_str, '%Y-%m-%d').date()
                        try:
                            rawmaterial = RawMaterial.objects.get(sku=sku_rawmaterial)
                            if rawmaterial.name != Nama_rawmaterial:
                                messages.warning(request, f"Raw material name '{Nama_rawmaterial}' does not match database name '{rawmaterial.name}' for SKU {sku_rawmaterial}, skipping row.")
                                continue
                        except ObjectDoesNotExist:
                            messages.warning(request, f"SKU {sku_rawmaterial} not found in raw materials, skipping row.")
                            continue
                        penerimaan, created = PenerimaanRawMaterial.objects.get_or_create(
                            id_penerimaan=id_penerimaan,
                            defaults={'id_supplier': Supplier, 'tanggal_terima': tanggal_terima, 'keterangan': ''}
                        )
                        PenerimaanRawMaterialDetail.objects.create(
                            id_penerimaan=penerimaan,
                            tanggal_terima=tanggal_terima,
                            Supplier=Supplier,
                            BRAND=BRAND,
                            id_rawmaterial=rawmaterial,
                            Nama_rawmaterial=Nama_rawmaterial,
                            quantity=int(quantity),
                            satuan=satuan,
                            harga_satuan=float(harga_satuan),
                            harga_total=int(quantity) * float(harga_satuan),
                            kategori=kategori
                        )
                        created_count += 1
                    except Exception as e:
                        messages.error(request, f"Error processing row for {id_penerimaan}: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} details created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('penerimaan_rawmaterial_list')
    else:
        form = BulkUploadForm()
    return render(request, 'inventory_product/penerimaan_rawmaterial_bulk_form.html', {'form': form, 'title': 'Bulk Upload Penerimaan Raw Material Details'})

def penerimaan_rawmaterial_detail_list(request, pk):
    penerimaan = get_object_or_404(PenerimaanRawMaterial, pk=pk)
    details = PenerimaanRawMaterialDetail.objects.filter(id_penerimaan=penerimaan)
    return render(request, 'inventory_product/penerimaan_rawmaterial_detail_list.html', {'penerimaan': penerimaan, 'details': details})

def penerimaan_rawmaterial_detail_create(request, pk):
    penerimaan = get_object_or_404(PenerimaanRawMaterial, pk=pk)
    if request.method == 'POST':
        tanggal_terima = request.POST.get('tanggal_terima')
        Supplier = request.POST.get('Supplier')
        BRAND = request.POST.get('BRAND')
        id_rawmaterial_id = request.POST.get('id_rawmaterial')
        Nama_rawmaterial = request.POST.get('Nama_rawmaterial')
        quantity = int(request.POST.get('quantity'))
        satuan = request.POST.get('satuan')
        harga_satuan = float(request.POST.get('harga_satuan'))
        harga_total = quantity * harga_satuan
        kategori = request.POST.get('kategori')
        rawmaterial = get_object_or_404(RawMaterial, pk=id_rawmaterial_id)
        PenerimaanRawMaterialDetail.objects.create(
            id_penerimaan=penerimaan,
            tanggal_terima=tanggal_terima,
            Supplier=Supplier,
            BRAND=BRAND,
            id_rawmaterial=rawmaterial,
            Nama_rawmaterial=Nama_rawmaterial,
            quantity=quantity,
            satuan=satuan,
            harga_satuan=harga_satuan,
            harga_total=harga_total,
            kategori=kategori
        )
        messages.success(request, 'Detail added successfully.')
        return redirect('penerimaan_rawmaterial_detail_list', pk=pk)
    rawmaterials = RawMaterial.objects.all()
    return render(request, 'inventory_product/penerimaan_rawmaterial_detail_form.html', {'penerimaan': penerimaan, 'rawmaterials': rawmaterials, 'title': 'Add Detail'})

# Similar for Keluar Raw Material
def keluar_rawmaterial_list(request):
    keluars = KeluarRawMaterial.objects.all()
    return render(request, 'inventory_product/keluar_rawmaterial_list.html', {'keluars': keluars})

def keluar_rawmaterial_create(request):
    if request.method == 'POST':
        id_keluar = request.POST.get('id_keluar')
        tanggal_keluar = request.POST.get('tanggal_keluar')
        keterangan = request.POST.get('keterangan')
        KeluarRawMaterial.objects.create(
            id_keluar=id_keluar,
            tanggal_keluar=tanggal_keluar,
            keterangan=keterangan
        )
        messages.success(request, 'Keluar Raw Material created successfully.')
        return redirect('keluar_rawmaterial_list')
    return render(request, 'inventory_product/keluar_rawmaterial_form.html', {'title': 'Create Keluar Raw Material'})

def keluar_rawmaterial_update(request, pk):
    keluar = get_object_or_404(KeluarRawMaterial, pk=pk)
    if request.method == 'POST':
        keluar.tanggal_keluar = request.POST.get('tanggal_keluar')
        keluar.keterangan = request.POST.get('keterangan')
        keluar.save()
        messages.success(request, 'Keluar Raw Material updated successfully.')
        return redirect('keluar_rawmaterial_list')
    return render(request, 'inventory_product/keluar_rawmaterial_form.html', {'keluar': keluar, 'title': 'Update Keluar Raw Material'})

def keluar_rawmaterial_delete(request, pk):
    keluar = get_object_or_404(KeluarRawMaterial, pk=pk)
    if request.method == 'POST':
        keluar.delete()
        messages.success(request, 'Keluar Raw Material deleted successfully.')
        return redirect('keluar_rawmaterial_list')
    return render(request, 'inventory_product/keluar_rawmaterial_confirm_delete.html', {'keluar': keluar})

def keluar_rawmaterial_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 8:
                        continue
                    id_keluar, tanggal_keluar_str, BRAND, sku_rawmaterial, Nama_rawmaterial, quantity, harga_satuan, kategori = row[:8]
                    try:
                        tanggal_keluar = datetime.strptime(tanggal_keluar_str, '%Y-%m-%d').date()
                        try:
                            rawmaterial = RawMaterial.objects.get(sku=sku_rawmaterial)
                            if rawmaterial.name != Nama_rawmaterial:
                                messages.warning(request, f"Raw material name '{Nama_rawmaterial}' does not match database name '{rawmaterial.name}' for SKU {sku_rawmaterial}, skipping row.")
                                continue
                        except ObjectDoesNotExist:
                            messages.warning(request, f"SKU {sku_rawmaterial} not found in raw materials, skipping row.")
                            continue
                        keluar, created = KeluarRawMaterial.objects.get_or_create(
                            id_keluar=id_keluar,
                            defaults={'tanggal_keluar': tanggal_keluar, 'keterangan': ''}
                        )
                        KeluarRawMaterialDetail.objects.create(
                            id_keluar=keluar,
                            tanggal_keluar=tanggal_keluar,
                            BRAND=BRAND,
                            id_rawmaterial=rawmaterial,
                            Nama_rawmaterial=Nama_rawmaterial,
                            quantity=int(quantity),
                            harga_satuan=float(harga_satuan),
                            harga_total=int(quantity) * float(harga_satuan),
                            kategori=kategori
                        )
                        created_count += 1
                    except Exception as e:
                        messages.error(request, f"Error processing row for {id_keluar}: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} details created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('keluar_rawmaterial_list')
    else:
        form = BulkUploadForm()
    return render(request, 'inventory_product/keluar_rawmaterial_bulk_form.html', {'form': form, 'title': 'Bulk Upload Keluar Raw Material Details'})

def keluar_rawmaterial_detail_list(request, pk):
    keluar = get_object_or_404(KeluarRawMaterial, pk=pk)
    details = KeluarRawMaterialDetail.objects.filter(id_keluar=keluar)
    return render(request, 'inventory_product/keluar_rawmaterial_detail_list.html', {'keluar': keluar, 'details': details})

def keluar_rawmaterial_detail_create(request, pk):
    keluar = get_object_or_404(KeluarRawMaterial, pk=pk)
    if request.method == 'POST':
        tanggal_keluar = request.POST.get('tanggal_keluar')
        BRAND = request.POST.get('BRAND')
        id_rawmaterial_id = request.POST.get('id_rawmaterial')
        Nama_rawmaterial = request.POST.get('Nama_rawmaterial')
        quantity = int(request.POST.get('quantity'))
        harga_satuan = float(request.POST.get('harga_satuan'))
        harga_total = quantity * harga_satuan
        kategori = request.POST.get('kategori')
        rawmaterial = get_object_or_404(RawMaterial, pk=id_rawmaterial_id)
        KeluarRawMaterialDetail.objects.create(
            id_keluar=keluar,
            tanggal_keluar=tanggal_keluar,
            BRAND=BRAND,
            id_rawmaterial=rawmaterial,
            Nama_rawmaterial=Nama_rawmaterial,
            quantity=quantity,
            harga_satuan=harga_satuan,
            harga_total=harga_total,
            kategori=kategori
        )
        messages.success(request, 'Detail added successfully.')
        return redirect('keluar_rawmaterial_detail_list', pk=pk)
    rawmaterials = RawMaterial.objects.all()
    return render(request, 'inventory_product/keluar_rawmaterial_detail_form.html', {'keluar': keluar, 'rawmaterials': rawmaterials, 'title': 'Add Detail'})

# Download templates for raw materials
def download_penerimaan_rawmaterial_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Penerimaan Raw Material Details"
    headers = ['ID Penerimaan', 'Tanggal Terima', 'Supplier', 'BRAND', 'SKU Raw Material', 'Nama Raw Material', 'Quantity', 'Satuan', 'Harga Satuan', 'Kategori']
    ws.append(headers)
    ws.append(['PRM001', '2025-01-01', 'Supplier A', 'Brand X', 'RM001', 'Sample Raw Material', '10', 'kg', '15.00', 'Category A'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=penerimaan_rawmaterial_template.xlsx'
    wb.save(response)
    return response

def download_keluar_rawmaterial_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Keluar Raw Material Details"
    headers = ['ID Keluar', 'Tanggal Keluar', 'BRAND', 'SKU Raw Material', 'Nama Raw Material', 'Quantity', 'Harga Satuan', 'Kategori']
    ws.append(headers)
    ws.append(['KLR001', '2025-01-01', 'Brand X', 'RM001', 'Sample Raw Material', '5', '20.00', 'Category A'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=keluar_rawmaterial_template.xlsx'
    wb.save(response)
    return response

# Download templates
def download_penerimaan_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Penerimaan Details"
    headers = ['ID Penerimaan', 'Tanggal Terima', 'Supplier', 'BRAND', 'SKU Product', 'Nama Product', 'Quantity', 'Satuan', 'Harga Satuan', 'Kategori']
    ws.append(headers)
    ws.append(['PNR001', '2025-01-01', 'Supplier A', 'Brand X', 'PROD001', 'Sample Product', '10', 'pcs', '15.00', 'Category A'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=penerimaan_template.xlsx'
    wb.save(response)
    return response

def download_keluar_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Keluar Details"
    headers = ['ID Keluar', 'Tanggal Keluar', 'BRAND', 'SKU Product', 'Nama Product', 'Quantity', 'Harga Satuan', 'Kategori']
    ws.append(headers)
    ws.append(['KLR001', '2025-01-01', 'Brand X', 'PROD001', 'Sample Product', '5', '20.00', 'Category A'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=keluar_template.xlsx'
    wb.save(response)
    return response

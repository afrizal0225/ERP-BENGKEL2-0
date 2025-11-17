from django.shortcuts import render, get_object_or_404, redirect
from .models import Product, RawMaterial
from .forms import ProductForm, RawMaterialForm, BulkUploadForm
from django.contrib import messages
import openpyxl
from datetime import datetime
from openpyxl import Workbook
from django.http import HttpResponse

def product_list(request):
    products = Product.objects.all()
    return render(request, 'product/product_list.html', {'products': products})

def product_create(request):
    if request.method == 'POST':
        form = ProductForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Product created successfully.')
            return redirect('product_list')
    else:
        form = ProductForm()
    return render(request, 'product/product_form.html', {'form': form, 'title': 'Create Product'})

def product_update(request, pk):
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        form = ProductForm(request.POST, instance=product)
        if form.is_valid():
            form.save()
            messages.success(request, 'Product updated successfully.')
            return redirect('product_list')
    else:
        form = ProductForm(instance=product)
    return render(request, 'product/product_form.html', {'form': form, 'title': 'Update Product'})

def product_delete(request, pk):
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        product.delete()
        messages.success(request, 'Product deleted successfully.')
        return redirect('product_list')
    return render(request, 'product/product_confirm_delete.html', {'product': product})

def product_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 8:
                        continue
                    sku, name, description, price, category, size, colour, release_date_str = row[:8]
                    try:
                        release_date = datetime.strptime(release_date_str, '%Y-%m-%d').date() if release_date_str else None
                        Product.objects.create(
                            sku=sku,
                            name=name,
                            description=description or '',
                            price=float(price) if price else 0,
                            category=category,
                            size=size or '',
                            colour=colour or '',
                            release_date=release_date,
                        )
                        created_count += 1
                    except Exception as e:
                        messages.error(request, f"Error creating product {name}: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} products created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('product_list')
    else:
        form = BulkUploadForm()
    return render(request, 'product/product_bulk_form.html', {'form': form, 'title': 'Bulk Upload Products'})

def rawmaterial_list(request):
    rawmaterials = RawMaterial.objects.all()
    return render(request, 'product/rawmaterial_list.html', {'rawmaterials': rawmaterials})

def rawmaterial_create(request):
    if request.method == 'POST':
        form = RawMaterialForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Raw Material created successfully.')
            return redirect('rawmaterial_list')
    else:
        form = RawMaterialForm()
    return render(request, 'product/rawmaterial_form.html', {'form': form, 'title': 'Create Raw Material'})

def rawmaterial_update(request, pk):
    rawmaterial = get_object_or_404(RawMaterial, pk=pk)
    if request.method == 'POST':
        form = RawMaterialForm(request.POST, instance=rawmaterial)
        if form.is_valid():
            form.save()
            messages.success(request, 'Raw Material updated successfully.')
            return redirect('rawmaterial_list')
    else:
        form = RawMaterialForm(instance=rawmaterial)
    return render(request, 'product/rawmaterial_form.html', {'form': form, 'title': 'Update Raw Material'})

def rawmaterial_delete(request, pk):
    rawmaterial = get_object_or_404(RawMaterial, pk=pk)
    if request.method == 'POST':
        rawmaterial.delete()
        messages.success(request, 'Raw Material deleted successfully.')
        return redirect('rawmaterial_list')
    return render(request, 'product/rawmaterial_confirm_delete.html', {'rawmaterial': rawmaterial})

def rawmaterial_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 6:
                        continue
                    sku, name, description, price, unit_of_measurement, category = row[:6]
                    try:
                        RawMaterial.objects.create(
                            sku=sku,
                            name=name,
                            description=description or '',
                            price=float(price) if price else 0,
                            unit_of_measurement=unit_of_measurement,
                            category=category,
                        )
                        created_count += 1
                    except Exception as e:
                        messages.error(request, f"Error creating raw material {name}: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} raw materials created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('rawmaterial_list')
    else:
        form = BulkUploadForm()
    return render(request, 'product/rawmaterial_bulk_form.html', {'form': form, 'title': 'Bulk Upload Raw Materials'})

def download_product_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    headers = ['SKU', 'Name', 'Description', 'Price', 'Category', 'Size', 'Colour', 'Release Date']
    ws.append(headers)
    # Sample row
    ws.append(['PROD001', 'Sample Shoe', 'High-quality leather shoe', '150.00', 'Footwear', '42', 'Black', '2025-01-01'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=product_template.xlsx'
    wb.save(response)
    return response

def items_list(request):
    products = Product.objects.all()
    rawmaterials = RawMaterial.objects.all()
    items = []
    for p in products:
        items.append({'type': 'Product', 'sku': p.sku, 'name': p.name, 'price': p.price, 'category': p.category, 'min_stock': p.minimum_stock, 'max_stock': p.maximum_stock, 'object': p, 'url_prefix': 'product'})
    for r in rawmaterials:
        items.append({'type': 'Raw Material', 'sku': r.sku, 'name': r.name, 'price': r.price, 'category': r.category, 'min_stock': r.minimum_stock, 'max_stock': r.maximum_stock, 'object': r, 'url_prefix': 'rawmaterial'})
    return render(request, 'product/items_list.html', {'items': items})

def download_rawmaterial_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Materials"
    headers = ['SKU', 'Name', 'Description', 'Price', 'Unit of Measurement', 'Category']
    ws.append(headers)
    # Sample row
    ws.append(['MAT001', 'Leather', 'Premium cow leather', '50.00', 'sq ft', 'Material'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=rawmaterial_template.xlsx'
    wb.save(response)
    return response

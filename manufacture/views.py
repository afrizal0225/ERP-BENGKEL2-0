from django.shortcuts import render, get_object_or_404, redirect
from .models import StasiunKerja, BOM, BOMDetail, ProductionOrder, ProductionOrderDetail, SuratPerintahKerja, SPKDetail, SPKOutput, ProductionProgress
from product.models import Product, RawMaterial
from django.db.models import Q
from .forms import StasiunKerjaForm, BOMForm, BOMDetailForm, BulkUploadForm, ProductionOrderForm, ProductionOrderDetailForm, SuratPerintahKerjaForm, ProductionProgressForm
from django.contrib import messages
import openpyxl
from datetime import datetime
from openpyxl import Workbook
from django.http import HttpResponse, JsonResponse
from django.db.models import Sum
from decimal import Decimal

def stasiunkerja_list(request):
    stasiunkers = StasiunKerja.objects.all()
    return render(request, 'manufacture/stasiunkerja_list.html', {'stasiunkers': stasiunkers})

def stasiunkerja_create(request):
    if request.method == 'POST':
        form = StasiunKerjaForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Stasiun Kerja created successfully.')
            return redirect('stasiunkerja_list')
    else:
        form = StasiunKerjaForm()
    return render(request, 'manufacture/stasiunkerja_form.html', {'form': form, 'title': 'Create Stasiun Kerja'})

def stasiunkerja_update(request, pk):
    stasiunker = get_object_or_404(StasiunKerja, pk=pk)
    if request.method == 'POST':
        form = StasiunKerjaForm(request.POST, instance=stasiunker)
        if form.is_valid():
            form.save()
            messages.success(request, 'Stasiun Kerja updated successfully.')
            return redirect('stasiunkerja_list')
    else:
        form = StasiunKerjaForm(instance=stasiunker)
    return render(request, 'manufacture/stasiunkerja_form.html', {'form': form, 'title': 'Update Stasiun Kerja'})

def stasiunkerja_delete(request, pk):
    stasiunker = get_object_or_404(StasiunKerja, pk=pk)
    if request.method == 'POST':
        stasiunker.delete()
        messages.success(request, 'Stasiun Kerja deleted successfully.')
        return redirect('stasiunkerja_list')
    return render(request, 'manufacture/stasiunkerja_confirm_delete.html', {'stasiunker': stasiunker})

def stasiunkerja_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 2:
                        continue
                    nama_stasiun_kerja, keterangan = row[:2]
                    try:
                        StasiunKerja.objects.create(
                            nama_stasiun_kerja=nama_stasiun_kerja,
                            keterangan=keterangan or '',
                        )
                        created_count += 1
                    except Exception as e:
                        messages.error(request, f"Error creating stasiun kerja {nama_stasiun_kerja}: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} stasiun kerja created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('stasiunkerja_list')
    else:
        form = BulkUploadForm()
    return render(request, 'manufacture/stasiunkerja_bulk_form.html', {'form': form, 'title': 'Bulk Upload Stasiun Kerja'})

def bom_list(request):
    boms = BOM.objects.all()
    return render(request, 'manufacture/bom_list.html', {'boms': boms})

def bom_create(request):
    if request.method == 'POST':
        form = BOMForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'BOM created successfully.')
            return redirect('bom_list')
    else:
        form = BOMForm()
    return render(request, 'manufacture/bom_form.html', {'form': form, 'title': 'Create BOM'})

def bom_update(request, pk):
    bom = get_object_or_404(BOM, pk=pk)
    details = BOMDetail.objects.filter(id_bom=bom)
    rawmaterials = RawMaterial.objects.all()
    stasiunkers = StasiunKerja.objects.all()
    return render(request, 'manufacture/bom_update.html', {
        'bom': bom,
        'details': details,
        'rawmaterials': rawmaterials,
        'stasiunkers': stasiunkers,
        'title': 'Update BOM Details'
    })

def bom_delete(request, pk):
    bom = get_object_or_404(BOM, pk=pk)
    if request.method == 'POST':
        bom.delete()
        messages.success(request, 'BOM deleted successfully.')
        return redirect('bom_list')
    return render(request, 'manufacture/bom_confirm_delete.html', {'bom': bom})

def bom_detail(request, pk):
    bom = get_object_or_404(BOM, pk=pk)
    details = BOMDetail.objects.filter(id_bom=bom)
    total_hpp = sum(detail.harga_totalrawmaterial for detail in details)
    return render(request, 'manufacture/bom_detail.html', {'bom': bom, 'details': details, 'total_hpp': total_hpp})

def bom_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 10:
                        continue
                    bom_id, product_sku, nama_product, version, rawmaterial_sku, nama_rawmaterial, stasiunkerja_id, nama_stasiunkerja, satuan, qty = row[:10]
                    try:
                        product = Product.objects.get(sku=product_sku)
                        bom, bom_created = BOM.objects.get_or_create(
                            id_bom=bom_id,
                            defaults={'id_product': product, 'version': version}
                        )
                        rawmaterial = RawMaterial.objects.get(sku=rawmaterial_sku)
                        stasiunkerja = StasiunKerja.objects.get(id_stasiun_kerja=stasiunkerja_id)
                        BOMDetail.objects.get_or_create(
                            id_bom=bom,
                            id_product=product,
                            version=version,
                            id_rawmaterial=rawmaterial,
                            defaults={
                                'nama_product': nama_product or '',
                                'nama_rawmaterial': nama_rawmaterial or '',
                                'id_stasiunkerja': stasiunkerja,
                                'nama_stasiunkerja': nama_stasiunkerja or '',
                                'satuan': satuan,
                                'qty': qty,
                                'harga_satuanrawmaterial': rawmaterial.price,
                                'harga_totalrawmaterial': rawmaterial.price * qty,
                            }
                        )
                        created_count += 1
                    except Product.DoesNotExist:
                        messages.error(request, f"Product with SKU {product_sku} not found")
                    except RawMaterial.DoesNotExist:
                        messages.error(request, f"Raw Material with SKU {rawmaterial_sku} not found")
                    except StasiunKerja.DoesNotExist:
                        messages.error(request, f"Stasiun Kerja with ID {stasiunkerja_id} not found")
                    except Exception as e:
                        messages.error(request, f"Error creating BOM Detail: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} BOM Details created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('bom_list')
    else:
        form = BulkUploadForm()
    return render(request, 'manufacture/bom_bulk_form.html', {'form': form, 'title': 'Bulk Upload BOM Details'})

def download_stasiunkerja_template(request):
    wb = Workbook()
def bomdetail_update(request, pk):
    bomdetail = get_object_or_404(BOMDetail, pk=pk)
    if request.method == 'POST':
        form = BOMDetailForm(request.POST, instance=bomdetail)
        if form.is_valid():
            form.save()
            messages.success(request, 'BOM Detail updated successfully.')
            return redirect('bom_update', pk=bomdetail.id_bom.pk)
    else:
        form = BOMDetailForm(instance=bomdetail)
    return render(request, 'manufacture/bomdetail_form.html', {'form': form, 'title': 'Update BOM Detail'})

def bomdetail_delete(request, pk):
    bomdetail = get_object_or_404(BOMDetail, pk=pk)
    bom_pk = bomdetail.id_bom.pk
    if request.method == 'POST':
        bomdetail.delete()
        messages.success(request, 'BOM Detail deleted successfully.')
        return redirect('bom_update', pk=bom_pk)
    return render(request, 'manufacture/bomdetail_confirm_delete.html', {'bomdetail': bomdetail})

def productionorder_list(request):
    productionorders = ProductionOrder.objects.all()
    return render(request, 'manufacture/productionorder_list.html', {'productionorders': productionorders})


def productionorder_update(request, pk):
    productionorder = get_object_or_404(ProductionOrder, pk=pk)
    if request.method == 'POST':
        form = ProductionOrderForm(request.POST, instance=productionorder)
        if form.is_valid():
            form.save()
            messages.success(request, 'Production Order updated successfully.')
            return redirect('productionorder_list')
    else:
        form = ProductionOrderForm(instance=productionorder)
    return render(request, 'manufacture/productionorder_form.html', {'form': form, 'title': 'Update Production Order'})

def productionorder_delete(request, pk):
    productionorder = get_object_or_404(ProductionOrder, pk=pk)
    if request.method == 'POST':
        productionorder.delete()
        messages.success(request, 'Production Order deleted successfully.')
        return redirect('productionorder_list')
    return render(request, 'manufacture/productionorder_confirm_delete.html', {'productionorder': productionorder})

def productionorder_detail(request, pk):
    productionorder = get_object_or_404(ProductionOrder, pk=pk)
    details = ProductionOrderDetail.objects.filter(id_po=productionorder)
    return render(request, 'manufacture/productionorder_detail.html', {'productionorder': productionorder, 'details': details})

def productionorderdetail_create(request, po_pk):
    productionorder = get_object_or_404(ProductionOrder, pk=po_pk)
    if request.method == 'POST':
        form = ProductionOrderDetailForm(request.POST)
        if form.is_valid():
            product = form.cleaned_data['id_product']
            # Check if product has BOM
            if not BOM.objects.filter(id_product=product).exists():
                messages.error(request, f"Cannot create Production Order Detail: Product '{product.name}' (SKU: {product.sku}) does not have a BOM. Production Orders can only be created for products with existing BOMs.")
                return render(request, 'manufacture/productionorderdetail_form.html', {'form': form, 'productionorder': productionorder, 'title': 'Create Production Order Detail'})

            detail = form.save(commit=False)
            detail.id_po = productionorder
            detail.save()
            messages.success(request, 'Production Order Detail created successfully.')
            return redirect('productionorder_detail', pk=po_pk)
    else:
        form = ProductionOrderDetailForm()
    return render(request, 'manufacture/productionorderdetail_form.html', {'form': form, 'productionorder': productionorder, 'title': 'Create Production Order Detail'})

def productionorderdetail_update(request, pk):
    detail = get_object_or_404(ProductionOrderDetail, pk=pk)
    if request.method == 'POST':
        form = ProductionOrderDetailForm(request.POST, instance=detail)
        if form.is_valid():
            form.save()
            messages.success(request, 'Production Order Detail updated successfully.')
            return redirect('productionorder_detail', pk=detail.id_po.pk)
    else:
        form = ProductionOrderDetailForm(instance=detail)
    return render(request, 'manufacture/productionorderdetail_form.html', {'form': form, 'productionorder': detail.id_po, 'title': 'Update Production Order Detail'})

def productionorderdetail_delete(request, pk):
    detail = get_object_or_404(ProductionOrderDetail, pk=pk)
    po_pk = detail.id_po.pk
    if request.method == 'POST':
        detail.delete()
        messages.success(request, 'Production Order Detail deleted successfully.')
        return redirect('productionorder_detail', pk=po_pk)
    return render(request, 'manufacture/productionorderdetail_confirm_delete.html', {'detail': detail})

def productionorder_bulk_create(request):
    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            try:
                wb = openpyxl.load_workbook(file)
                sheet = wb.active
                created_count = 0
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) < 9:
                        continue
                    po_id, tanggal_request, tanggal_selesai, product_sku, nama_product, size, warna, qty_produksi, keterangan = row[:9]
                    try:
                        # Parse dates
                        if isinstance(tanggal_request, str):
                            tanggal_request = datetime.strptime(tanggal_request, '%Y-%m-%d').date()
                        if isinstance(tanggal_selesai, str):
                            tanggal_selesai = datetime.strptime(tanggal_selesai, '%Y-%m-%d').date()

                        product = Product.objects.get(sku=product_sku)

                        # Check if product has BOM
                        if not BOM.objects.filter(id_product=product).exists():
                            messages.error(request, f"Product with SKU {product_sku} does not have a BOM. Production Orders can only be created for products with existing BOMs.")
                            continue

                        po, po_created = ProductionOrder.objects.get_or_create(
                            id_po=po_id,
                            defaults={
                                'tanggal_request': tanggal_request,
                                'tanggal_selesai': tanggal_selesai
                            }
                        )
                        ProductionOrderDetail.objects.get_or_create(
                            id_po=po,
                            id_product=product,
                            defaults={
                                'nama_product': nama_product or '',
                                'size': size or '',
                                'warna': warna or '',
                                'qty_produksi': qty_produksi,
                                'keterangan': keterangan or '',
                            }
                        )
                        created_count += 1
                    except Product.DoesNotExist:
                        messages.error(request, f"Product with SKU {product_sku} not found")
                    except Exception as e:
                        messages.error(request, f"Error creating Production Order Detail: {e}")
                messages.success(request, f'Bulk upload completed. {created_count} Production Order Details created.')
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
            return redirect('productionorder_list')
    else:
        form = BulkUploadForm()
    return render(request, 'manufacture/productionorder_bulk_form.html', {'form': form, 'title': 'Bulk Upload Production Orders'})

    ws = wb.active
    ws.title = "Stasiun Kerja"
    headers = ['Nama Stasiun Kerja', 'Keterangan']
    ws.append(headers)
    # Sample row
    ws.append(['Assembly Line 1', 'Main assembly station'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=stasiunkerja_template.xlsx'
    wb.save(response)
    return response

def download_bom_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    headers = ['BOM ID', 'Product SKU', 'Nama Product', 'Version', 'Raw Material SKU', 'Nama Raw Material', 'Stasiun Kerja ID', 'Nama Stasiun Kerja', 'Satuan', 'Qty']
    ws.append(headers)
    # Sample row
    ws.append(['BOM001', 'PROD001', 'Sample Product', '1.0', 'RAW001', 'Sample Raw Material', '1', 'Sample Stasiun Kerja', 'pcs', '10.5'])

    # Sheet for existing products
    ws_products = wb.create_sheet("Products_Exist")
    ws_products.append(['Product SKU', 'Nama Product'])
    for product in Product.objects.all():
        ws_products.append([product.sku, product.name])

    # Sheet for existing raw materials
    ws_rawmaterials = wb.create_sheet("RawMaterials_Exist")
    ws_rawmaterials.append(['Raw Material SKU', 'Nama Raw Material','Category'])
    for rawmaterial in RawMaterial.objects.all():
        ws_rawmaterials.append([rawmaterial.sku, rawmaterial.name, rawmaterial.category])
    # Sheet for existing stasiun kerja
    ws_stasiun = wb.create_sheet("StasiunKerja_Exist")
    ws_stasiun.append(['Stasiun Kerja ID', 'Nama Stasiun Kerja'])
    for stasiun in StasiunKerja.objects.all():
        ws_stasiun.append([stasiun.id_stasiun_kerja, stasiun.nama_stasiun_kerja])


    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=bom_template.xlsx'
    wb.save(response)
    return response

def download_productionorder_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Production Order"
    headers = ['PO ID', 'Tanggal Request', 'Tanggal Selesai', 'Product SKU', 'Nama Product', 'Size', 'Warna', 'Qty Produksi', 'Keterangan']
    ws.append(headers)
    # Sample row
    ws.append(['PO001', '2024-01-15', '2024-01-20', 'PROD001', 'Sample Product', 'M', 'Red', '100', 'Urgent order'])

    # Sheet for existing products that have BOMs
    ws_products = wb.create_sheet("Products_Exist")
    ws_products.append(['Product SKU', 'Nama Product', 'Size', 'Colour', 'BOM Version'])
    # Only include products that have BOMs
    products_with_bom = Product.objects.filter(bom__isnull=False).distinct()
    for product in products_with_bom:
        # Get the latest BOM version for this product
        latest_bom = BOM.objects.filter(id_product=product).order_by('-version').first()
        bom_version = latest_bom.version if latest_bom else ''
        ws_products.append([product.sku, product.name, product.size, product.colour, bom_version])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=productionorder_template.xlsx'
    wb.save(response)
    return response

def spk_list(request):
    spks = SuratPerintahKerja.objects.all()

    # Handle filtering
    id_spk_filter = request.GET.get('id_spk', '')
    id_po_filter = request.GET.get('id_po', '')
    tanggal_spk_filter = request.GET.get('tanggal_spk', '')
    status_filter = request.GET.get('status', '')

    if id_spk_filter:
        spks = spks.filter(id_spk__icontains=id_spk_filter)
    if id_po_filter:
        spks = spks.filter(id_po__id_po__icontains=id_po_filter)
    if tanggal_spk_filter:
        spks = spks.filter(tanggal_spk=tanggal_spk_filter)
    if status_filter:
        spks = spks.filter(status=status_filter)

    context = {
        'spks': spks,
        'id_spk_filter': id_spk_filter,
        'id_po_filter': id_po_filter,
        'tanggal_spk_filter': tanggal_spk_filter,
        'status_filter': status_filter,
    }
    return render(request, 'manufacture/spk_list.html', context)

def spk_create(request):
    selected_po = request.POST.get('id_po') or request.GET.get('po')
    po_details = None
    if selected_po:
        try:
            po = ProductionOrder.objects.get(id_po=selected_po)
            po_details = ProductionOrderDetail.objects.filter(id_po=po, qty_remaining__gt=0)
        except ProductionOrder.DoesNotExist:
            pass

    if request.method == 'POST':
        form = SuratPerintahKerjaForm(request.POST)
        if 'allocations' in request.POST:
            if form.is_valid():
                spk = form.save()
                # Process allocations
                allocations = {}
                for key, value in request.POST.items():
                    if key.startswith('allocated_'):
                        po_detail_id = key.split('_')[1]
                        try:
                            qty = Decimal(value)
                            if qty > 0:
                                allocations[int(po_detail_id)] = qty
                        except ValueError:
                            pass

                # Validate allocations
                for po_detail_id, qty in allocations.items():
                    try:
                        po_detail = ProductionOrderDetail.objects.get(pk=po_detail_id)
                        if qty > po_detail.qty_remaining:
                            messages.error(request, f'Allocated quantity {qty} exceeds remaining {po_detail.qty_remaining} for {po_detail.nama_product}')
                            return redirect('spk_create')
                    except ProductionOrderDetail.DoesNotExist:
                        messages.error(request, 'Invalid PO detail')
                        return redirect('spk_create')

                # Update remaining and create SPK outputs
                for po_detail_id, qty in allocations.items():
                    po_detail = ProductionOrderDetail.objects.get(pk=po_detail_id)
                    po_detail.qty_remaining -= qty
                    po_detail.save()

                    # Create SPKOutput
                    SPKOutput.objects.create(
                        id_spk=spk,
                        id_po_detail=po_detail,
                        id_product=po_detail.id_product,
                        nama_product=po_detail.nama_product,
                        qty_output=qty
                    )

                # Generate SPK details based on allocations
                generate_spk_details(spk)
                messages.success(request, 'SPK created successfully with allocated quantities.')
                return redirect('spk_detail', pk=spk.pk)
    else:
        initial = {}
        if selected_po:
            try:
                po = ProductionOrder.objects.get(id_po=selected_po)
                initial['id_po'] = po
            except ProductionOrder.DoesNotExist:
                pass
        form = SuratPerintahKerjaForm(initial=initial)

    return render(request, 'manufacture/spk_form.html', {'form': form, 'title': 'Create SPK', 'po_details': po_details})

def spk_detail(request, pk):
    spk = get_object_or_404(SuratPerintahKerja, pk=pk)
    details = SPKDetail.objects.filter(id_spk=spk)
    outputs = SPKOutput.objects.filter(id_spk=spk)
    return render(request, 'manufacture/spk_detail.html', {'spk': spk, 'details': details, 'outputs': outputs})

def spk_delete(request, pk):
    spk = get_object_or_404(SuratPerintahKerja, pk=pk)
    if request.method == 'POST':
        spk.delete()
        messages.success(request, 'SPK deleted successfully.')
        return redirect('spk_list')
    return render(request, 'manufacture/spk_confirm_delete.html', {'spk': spk})

def spk_approve(request, pk):
    spk = get_object_or_404(SuratPerintahKerja, pk=pk)
    if spk.status == 'Draft':
        spk.status = 'Approved'
        spk.save()
        messages.success(request, f'SPK {spk.id_spk} has been approved.')
    else:
        messages.warning(request, f'SPK {spk.id_spk} is already approved.')
    return redirect('spk_detail', pk=pk)

def download_spk(request, pk):
    spk = get_object_or_404(SuratPerintahKerja, pk=pk)
    wb = Workbook()

    # Get all stations for this SPK
    stations = StasiunKerja.objects.filter(
        id_stasiun_kerja__in=SPKDetail.objects.filter(id_spk=spk).values('id_stasiunkerja')
    ).distinct()

    for station in stations:
        # Create sheet for each station
        ws = wb.create_sheet(title=station.nama_stasiun_kerja[:31])  # Excel sheet name limit

        # SPK Information
        ws.append(['SPK Information'])
        ws.append(['ID SPK', spk.id_spk])
        ws.append(['ID PO', str(spk.id_po)])
        ws.append(['Tanggal SPK', str(spk.tanggal_spk)])
        ws.append(['Status', spk.status])
        ws.append(['Keterangan', spk.keterangan])
        ws.append([])  # Empty row

        # Raw Material Requirements for this station
        ws.append(['Kebutuhan Bahan Baku'])
        ws.append(['Bahan Baku', 'Qty Kebutuhan', 'Satuan'])

        station_details = SPKDetail.objects.filter(id_spk=spk, id_stasiunkerja=station)
        for detail in station_details:
            ws.append([detail.nama_rawmaterial, detail.qty_kebutuhan, detail.satuan])

        ws.append([])  # Empty row

        # Output Quantities for this station
        ws.append(['Jumlah Output'])
        ws.append(['Product', 'Qty Output'])

        # Since SPKOutput no longer has station, list all outputs (they are per product)
        outputs = SPKOutput.objects.filter(id_spk=spk)
        for output in outputs:
            ws.append([output.nama_product, output.qty_output])

    # Remove default sheet if it exists
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=SPK_{spk.id_spk}.xlsx'
    wb.save(response)
    return response

def generate_spk_details(spk):
    # Use SPKOutput to get allocated quantities
    spk_outputs = SPKOutput.objects.filter(id_spk=spk)

    # Dictionary to accumulate raw material needs per station
    raw_material_needs = {}

    for spk_output in spk_outputs:
        product = spk_output.id_product
        qty_allocated = spk_output.qty_output

        # Get BOM for the product
        bom = BOM.objects.filter(id_product=product).first()
        if not bom:
            continue

        bom_details = BOMDetail.objects.filter(id_bom=bom)

        for bom_detail in bom_details:
            station = bom_detail.id_stasiunkerja
            rawmaterial = bom_detail.id_rawmaterial
            qty_needed = bom_detail.qty * qty_allocated

            key = (station.id_stasiun_kerja, rawmaterial.sku)
            if key not in raw_material_needs:
                raw_material_needs[key] = {
                    'station': station,
                    'rawmaterial': rawmaterial,
                    'qty': 0,
                    'satuan': bom_detail.satuan
                }
            raw_material_needs[key]['qty'] += qty_needed


    # Create SPKDetail instances
    for key, data in raw_material_needs.items():
        SPKDetail.objects.create(
            id_spk=spk,
            id_stasiunkerja=data['station'],
            nama_stasiunkerja=data['station'].nama_stasiun_kerja,
            id_rawmaterial=data['rawmaterial'],
            nama_rawmaterial=data['rawmaterial'].name,
            qty_kebutuhan=data['qty'],
            satuan=data['satuan']
        )

def productionprogress_list(request):
    progresses = ProductionProgress.objects.all()

    # Handle filtering
    id_spk_filter = request.GET.get('id_spk', '')
    id_stasiunkerja_filter = request.GET.get('id_stasiunkerja', '')
    tanggal_mulai_filter = request.GET.get('tanggal_mulai', '')
    tanggal_selesai_filter = request.GET.get('tanggal_selesai', '')

    if id_spk_filter:
        progresses = progresses.filter(id_spk__id_spk__icontains=id_spk_filter)
    if id_stasiunkerja_filter:
        progresses = progresses.filter(id_stasiunkerja__id_stasiun_kerja__icontains=id_stasiunkerja_filter)
    if tanggal_mulai_filter:
        progresses = progresses.filter(tanggal_mulai=tanggal_mulai_filter)
    if tanggal_selesai_filter:
        progresses = progresses.filter(tanggal_selesai=tanggal_selesai_filter)

    context = {
        'progresses': progresses,
        'id_spk_filter': id_spk_filter,
        'id_stasiunkerja_filter': id_stasiunkerja_filter,
        'tanggal_mulai_filter': tanggal_mulai_filter,
        'tanggal_selesai_filter': tanggal_selesai_filter,
    }
    return render(request, 'manufacture/productionprogress_list.html', context)

def productionprogress_create(request):
    if request.method == 'POST':
        spk_id = request.POST.get('id_spk')
        station_id = request.POST.get('id_stasiunkerja')
        products_data = []

        # Parse products data
        index = 0
        while f'products[{index}][id_product]' in request.POST:
            product_data = {
                'id_product': request.POST.get(f'products[{index}][id_product]'),
                'tanggal_mulai': request.POST.get(f'products[{index}][tanggal_mulai]'),
                'tanggal_selesai': request.POST.get(f'products[{index}][tanggal_selesai]'),
                'qty_selesai': request.POST.get(f'products[{index}][qty_selesai]'),
            }
            products_data.append(product_data)
            index += 1

        if spk_id and station_id and products_data:
            try:
                spk = SuratPerintahKerja.objects.get(id_spk=spk_id)
                station = StasiunKerja.objects.get(id_stasiun_kerja=station_id)
                created_count = 0
                for data in products_data:
                    if data['qty_selesai'] and float(data['qty_selesai']) > 0:
                        product = Product.objects.get(sku=data['id_product'])
                        ProductionProgress.objects.create(
                            id_spk=spk,
                            id_stasiunkerja=station,
                            id_product=product,
                            nama_product=product.name,
                            tanggal_mulai=data['tanggal_mulai'],
                            tanggal_selesai=data['tanggal_selesai'],
                            qty_selesai=data['qty_selesai']
                        )
                        created_count += 1
                messages.success(request, f'{created_count} Production Progress entries created successfully.')
                return redirect('productionprogress_list')
            except Exception as e:
                messages.error(request, f'Error creating progress: {e}')
        else:
            messages.error(request, 'Invalid data submitted.')

    form = ProductionProgressForm()
    return render(request, 'manufacture/productionprogress_form.html', {'form': form, 'title': 'Create Production Progress'})

def productionprogress_update(request, pk):
    progress = get_object_or_404(ProductionProgress, pk=pk)
    if request.method == 'POST':
        form = ProductionProgressForm(request.POST, instance=progress)
        if form.is_valid():
            progress = form.save(commit=False)
            progress.nama_product = progress.id_product.name
            progress.save()
            messages.success(request, 'Production Progress updated successfully.')
            return redirect('productionprogress_list')
    else:
        form = ProductionProgressForm(instance=progress)
    return render(request, 'manufacture/productionprogress_form.html', {'form': form, 'title': 'Update Production Progress'})

def productionprogress_delete(request, pk):
    progress = get_object_or_404(ProductionProgress, pk=pk)
    if request.method == 'POST':
        progress.delete()
        messages.success(request, 'Production Progress deleted successfully.')
        return redirect('productionprogress_list')
    return render(request, 'manufacture/productionprogress_confirm_delete.html', {'progress': progress})

def manufacture_dashboard(request):
    # Get POs with remaining quantities (calculated as total PO output - total approved SPK outputs)
    pos_with_remaining = []
    all_pos = ProductionOrder.objects.all()
    for po in all_pos:
        total_po_output = po.productionorderdetail_set.aggregate(
            total=Sum('qty_produksi')
        )['total'] or 0

        total_spk_approved = SPKOutput.objects.filter(
            id_spk__id_po=po,
            id_spk__status='Approved'
        ).aggregate(total=Sum('qty_output'))['total'] or 0

        remaining = total_po_output - total_spk_approved
        if remaining > 0:
            po.total_remaining = remaining
            pos_with_remaining.append(po)

    # Get SPK progress per station
    spk_progress = ProductionProgress.objects.values(
        'id_spk__id_spk',
        'id_spk__id_po__id_po',
        'id_stasiunkerja__nama_stasiun_kerja',
        'id_stasiunkerja'
    ).annotate(
        total_qty_selesai=Sum('qty_selesai')
    ).order_by('id_spk__id_spk', 'id_stasiunkerja__nama_stasiun_kerja')

    # Calculate progress percentage for each
    progress_data = []
    for progress in spk_progress:
        spk_id = progress['id_spk__id_spk']
        station_id = progress['id_stasiunkerja']
        total_selesai = progress['total_qty_selesai'] or 0

        # Get total allocated for this SPK and station: sum of SPKOutput.qty_output
        # for products that have this station in their BOM
        products_with_station = BOMDetail.objects.filter(
            id_stasiunkerja=station_id
        ).values_list('id_product', flat=True).distinct()

        total_allocated = SPKOutput.objects.filter(
            id_spk__id_spk=spk_id,
            id_product__in=products_with_station
        ).aggregate(total=Sum('qty_output'))['total'] or 0

        percentage = (total_selesai / total_allocated * 100) if total_allocated > 0 else 0

        progress_data.append({
            'spk_id': spk_id,
            'po_id': progress['id_spk__id_po__id_po'],
            'station_name': progress['id_stasiunkerja__nama_stasiun_kerja'],
            'total_qty_selesai': total_selesai,
            'total_allocated': total_allocated,
            'percentage': round(percentage, 2)
        })

    # Get SPK details for context
    spks = SuratPerintahKerja.objects.filter(status='Approved').select_related('id_po')

    context = {
        'pos_with_remaining': pos_with_remaining,
        'spk_progress': progress_data,
        'spks': spks,
    }
    return render(request, 'manufacture/dashboard.html', context)

def get_products_for_spk(request):
    spk_id = request.GET.get('spk_id')
    if spk_id:
        try:
            spk = SuratPerintahKerja.objects.get(id_spk=spk_id)
            products = SPKOutput.objects.filter(id_spk=spk).select_related('id_product')
            data = [{'id': p.id_product.pk, 'sku': p.id_product.sku, 'name': p.id_product.name} for p in products]
            return JsonResponse({'products': data})
        except SuratPerintahKerja.DoesNotExist:
            return JsonResponse({'error': 'SPK not found'}, status=404)
    return JsonResponse({'error': 'No SPK ID provided'}, status=400)
